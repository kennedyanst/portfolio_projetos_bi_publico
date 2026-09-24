from datetime import datetime, timezone
from html import escape
from io import BytesIO
from pathlib import Path
import re
import unicodedata

import pandas as pd
import reportlab
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import CondPageBreak, HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


COL_DATA = "Data de Inicio"
COLUNAS_NUMERICAS = ("Projetos Contratados", "TRL - Inicial", "TRL - Final")
FONTE_DADOS = "dados/projetos_embrapii.csv"
VERDE = colors.HexColor("#125C49")
VERDE_CLARO = colors.HexColor("#EAF2EE")
CINZA = colors.HexColor("#52645C")


def preparar_dados(caminho: Path) -> pd.DataFrame:
    dados = pd.read_csv(
        caminho, sep=";", encoding="utf-8-sig", dtype="string",
        keep_default_na=False, na_values=[""],
    )
    dados[COL_DATA] = pd.to_datetime(dados[COL_DATA], format="%d/%m/%Y", errors="raise")
    for coluna in COLUNAS_NUMERICAS:
        dados[coluna] = pd.to_numeric(dados[coluna], errors="raise").astype("Int64")
    return dados


def filtrar_dados(dados: pd.DataFrame, filtros: dict, periodo: tuple | None) -> pd.DataFrame:
    resultado = dados
    if periodo:
        inicio, fim = periodo
        resultado = resultado.loc[resultado[COL_DATA].dt.date.between(inicio, fim)]
    for coluna, valores in filtros.items():
        if valores:
            resultado = resultado.loc[resultado[coluna].isin(valores)]
    return resultado


def opcoes_disponiveis(dados: pd.DataFrame, coluna: str, filtros: dict, periodo: tuple | None) -> list:
    outros_filtros = {nome: valores for nome, valores in filtros.items() if nome != coluna}
    segmentado = filtrar_dados(dados, outros_filtros, periodo)
    return segmentado[coluna].dropna().unique().tolist()


def descrever_filtros(filtros: dict, periodo: tuple | None = None) -> list[str]:
    descricao = []
    if periodo:
        inicio, fim = periodo
        descricao.append(f"{COL_DATA}: {inicio:%d/%m/%Y} a {fim:%d/%m/%Y}")
    for coluna, valores in filtros.items():
        if valores:
            descricao.append(f"{coluna}: {', '.join(str(valor) for valor in valores)}")
    return descricao


def gerar_excel(dados: pd.DataFrame) -> bytes:
    arquivo = BytesIO()
    planilha = dados.copy()
    for coluna in planilha.select_dtypes(include="string").columns:
        planilha[coluna] = planilha[coluna].str.replace(ILLEGAL_CHARACTERS_RE, " ", regex=True)
    with pd.ExcelWriter(arquivo, engine="openpyxl", datetime_format="DD/MM/YYYY") as writer:
        planilha.to_excel(writer, sheet_name="Projetos", index=False)
        aba = writer.sheets["Projetos"]
        aba.freeze_panes = "C2"
        aba.auto_filter.ref = aba.dimensions
        for celula in aba[1]:
            celula.fill = PatternFill("solid", fgColor="125C49")
            celula.font = Font(color="FFFFFF", bold=True)
            celula.alignment = Alignment(wrap_text=True, vertical="center")
        aba.row_dimensions[1].height = 35
        for indice, coluna in enumerate(dados.columns, start=1):
            largura = 55 if indice <= 2 else min(42, max(16, len(coluna) + 3))
            aba.column_dimensions[aba.cell(1, indice).column_letter].width = largura
    return arquivo.getvalue()


def _registrar_fontes() -> None:
    pasta = Path(reportlab.__file__).parent / "fonts"
    if "Vera" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("Vera", str(pasta / "Vera.ttf")))
        pdfmetrics.registerFont(TTFont("Vera-Bold", str(pasta / "VeraBd.ttf")))
        pdfmetrics.registerFontFamily("Vera", normal="Vera", bold="Vera-Bold")


def _texto(valor) -> str:
    if pd.isna(valor):
        return "—"
    if isinstance(valor, (pd.Timestamp, datetime)):
        return valor.strftime("%d/%m/%Y")
    texto = unicodedata.normalize("NFC", str(valor))
    texto = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", texto)
    # A fonte embarcada no ReportLab não contém alguns símbolos da base.
    substituicoes = {
        "●": "•", "₂": "2", "₅": "5", "₄": "4", "₀": "0", "₇": "7",
        "⁺": "+", "‑": "-", "\u200b": "", "\u202f": " ",
        "β": "beta", "Δ": "Delta", "→": "->", "\u0307": "",
    }
    return "".join(substituicoes.get(caractere, caractere) for caractere in texto)


def _paragrafo(valor, estilo: ParagraphStyle) -> Paragraph:
    texto = escape(_texto(valor)).replace("\r\n", "\n").replace("\r", "\n")
    return Paragraph(texto.replace("\n", "<br/>"), estilo)


def gerar_pdf(
    dados: pd.DataFrame,
    logo: Path | None = None,
    filtros: dict | None = None,
    periodo: tuple | None = None,
    origem_app: str | None = None,
    extraido_em: datetime | None = None,
) -> bytes:
    _registrar_fontes()
    extraido_em = extraido_em or datetime.now(timezone.utc)
    if extraido_em.tzinfo is None:
        extraido_em = extraido_em.replace(tzinfo=timezone.utc)
    instante = extraido_em.strftime("%d/%m/%Y às %H:%M:%S %Z")
    criterios = descrever_filtros(filtros or {}, periodo)
    arquivo = BytesIO()
    largura, altura = A4
    margem = 43
    largura_util = largura - 2 * margem
    documento = SimpleDocTemplate(
        arquivo,
        pagesize=A4,
        leftMargin=margem,
        rightMargin=margem,
        topMargin=72,
        bottomMargin=67,
        title="Projetos Embrapii — relatório de extração",
        author="Consulta de projetos Embrapii",
    )
    estilos = {
        "titulo": ParagraphStyle("titulo", fontName="Vera-Bold", fontSize=17, leading=23, textColor=VERDE, spaceAfter=8, splitLongWords=1),
        "subtitulo": ParagraphStyle("subtitulo", fontName="Vera-Bold", fontSize=9, leading=14, textColor=VERDE),
        "corpo": ParagraphStyle("corpo", fontName="Vera", fontSize=8.5, leading=13, textColor=colors.HexColor("#26352E"), splitLongWords=1),
        "pequeno": ParagraphStyle("pequeno", fontName="Vera", fontSize=7.4, leading=11, textColor=CINZA, splitLongWords=1),
        "rotulo": ParagraphStyle("rotulo", fontName="Vera-Bold", fontSize=7, leading=10, textColor=VERDE, splitLongWords=1),
        "registro": ParagraphStyle("registro", fontName="Vera-Bold", fontSize=10, leading=15, textColor=VERDE, splitLongWords=1),
    }

    def desenhar_pagina(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(VERDE)
        canvas.setLineWidth(1)
        canvas.line(margem, altura - 49, largura - margem, altura - 49)
        if logo is not None and logo.is_file():
            imagem = ImageReader(str(logo))
            w, h = imagem.getSize()
            escala = min(96 / w, 27 / h)
            canvas.drawImage(imagem, margem, altura - 42, width=w * escala, height=h * escala, mask="auto")
        canvas.setFont("Vera-Bold", 8)
        canvas.setFillColor(VERDE)
        canvas.drawRightString(largura - margem, altura - 34, "RELATÓRIO DE PROJETOS")
        canvas.setStrokeColor(colors.HexColor("#D4DFD8"))
        canvas.line(margem, 52, largura - margem, 52)
        canvas.setFont("Vera", 7)
        canvas.setFillColor(CINZA)
        canvas.drawString(margem, 39, f"Fonte: {FONTE_DADOS}  |  Extraído em {instante}")
        canvas.drawRightString(largura - margem, 39, f"Página {doc.page}")
        canvas.restoreState()

    historia = [
        _paragrafo("Projetos Embrapii | " + ("; ".join(criterios) if criterios else "sem filtros aplicados"), estilos["titulo"]),
        _paragrafo(f"{len(dados):,} projeto(s) encontrado(s)".replace(",", "."), estilos["subtitulo"]),
        Spacer(1, 9),
    ]
    origem = origem_app or "Aplicativo Streamlit (endereço indisponível)"
    for rotulo, valor in (("Fonte dos dados", FONTE_DADOS), ("Extraído de", origem), ("Data e hora da extração", instante)):
        historia.append(_paragrafo(f"{rotulo}: {valor}", estilos["pequeno"]))
    historia.extend([Spacer(1, 11), HRFlowable(width="100%", thickness=1, color=VERDE), Spacer(1, 12)])

    colunas = list(dados.columns)
    for indice, linha in enumerate(dados.itertuples(index=False, name=None), start=1):
        valores = dict(zip(colunas, linha))
        historia.append(CondPageBreak(90))
        historia.append(_paragrafo(f"PROJETO {indice:04d}  |  {colunas[0].upper()}", estilos["rotulo"]))
        historia.append(_paragrafo(valores[colunas[0]], estilos["registro"]))
        historia.append(Spacer(1, 5))
        historia.append(_paragrafo("DESCRIÇÃO", estilos["rotulo"]))
        historia.append(_paragrafo(valores[colunas[1]], estilos["corpo"]))
        historia.append(Spacer(1, 8))

        linhas = []
        detalhes = colunas[2:]
        for posicao in range(0, len(detalhes), 2):
            esquerda = detalhes[posicao]
            direita = detalhes[posicao + 1] if posicao + 1 < len(detalhes) else None
            linhas.append([
                _paragrafo(esquerda, estilos["rotulo"]),
                _paragrafo(valores[esquerda], estilos["corpo"]),
                _paragrafo(direita, estilos["rotulo"]) if direita else "",
                _paragrafo(valores[direita], estilos["corpo"]) if direita else "",
            ])
        tabela = Table(linhas, colWidths=[83, 166, 83, largura_util - 332], hAlign="LEFT", splitByRow=1)
        tabela.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("ROWBACKGROUNDS", (0, 0), (-1, -1), [VERDE_CLARO, colors.white]),
            ("LINEBELOW", (0, -1), (-1, -1), 0.5, colors.HexColor("#D4DFD8")),
        ]))
        historia.extend([tabela, Spacer(1, 14)])

    if dados.empty:
        historia.append(_paragrafo("Nenhum projeto corresponde aos filtros selecionados.", estilos["corpo"]))
    documento.build(historia, onFirstPage=desenhar_pagina, onLaterPages=desenhar_pagina)
    return arquivo.getvalue()
